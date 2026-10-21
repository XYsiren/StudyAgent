"""Learning acceptance tests use real SQL/Qdrant and deterministic model fixtures."""

import io
import json
import time
from datetime import timedelta
from types import SimpleNamespace

import pytest
from app.models import Chunk, Document, utcnow
from app.models_learning import Course, QuizAttempt, QuizItem, ReviewCard
from app.services.learning import today
from app.services.parsing import chunk_pages, parse_document
from app.services.pptx_parser import VisionPending
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy import select
from test_api import register, upload, wait_job


def create_course(client, headers, **kwargs):
    response = client.post(
        "/api/courses",
        headers=headers,
        json={"name": "深度学习", "exam_date": (today() + timedelta(days=30)).isoformat(), **kwargs},
    )
    assert response.status_code == 201, response.text
    return response.json()


def point(client, headers, course, **kwargs):
    return client.post(
        f"/api/courses/{course['id']}/knowledge-points",
        headers=headers,
        json={"title": "带薪年假", "chunk_ids": [], **kwargs},
    )


def test_course_required_fields_access_plan_and_delete(client):
    owner = register(client)
    assert client.post("/api/courses", headers=owner, json={"name": "缺考试日期"}).status_code == 422
    course = create_course(client, owner)
    assert course["exam_name"] == "深度学习" and course["weekly_minutes"] == 300
    p = client.get(f"/api/courses/{course['id']}/study-plan", headers=owner).json()
    assert p["tasks"] and p["tasks"][0]["task_type"] == "setup"
    assert sum(t["minutes"] for t in p["tasks"]) <= p["daily_budget_minutes"]
    first = p["tasks"][0]
    assert (
        client.post(f"/api/study-tasks/{first['id']}/complete", headers=owner, json={"status": "completed"}).status_code
        == 200
    )
    exam = (today() + timedelta(days=2)).isoformat()
    assert client.patch(f"/api/courses/{course['id']}", headers=owner, json={"exam_date": exam}).status_code == 200
    after = client.get(f"/api/courses/{course['id']}/study-plan", headers=owner).json()
    assert any(t["id"] == first["id"] and t["status"] == "completed" for t in after["tasks"])
    assert all(t["scheduled_for"] <= exam for t in after["tasks"] if t["status"] == "pending")
    stranger = register(client, "stranger@example.com")
    for path in ["", "/overview", "/materials", "/metrics/overview", "/study-plan", "/quiz", "/review/today"]:
        assert client.get(f"/api/courses/{course['id']}" + path, headers=stranger).status_code == 404
    assert client.delete(f"/api/courses/{course['id']}", headers=owner).status_code == 204
    assert client.get(f"/api/courses/{course['id']}/quiz", headers=owner).status_code == 404
    with client.app.state.database.session() as db:
        assert db.get(Course, course["id"]) is None


def test_past_exam_stops_plans_and_unconfirmed_points_do_not_enter_plan(client):
    headers = register(client)
    course = create_course(client, headers)
    doc = upload(client, headers, course["knowledge_base_id"])
    assert wait_job(client, headers, doc["job_id"])["status"] == "completed"
    suggested = client.post(f"/api/courses/{course['id']}/knowledge-points/suggest", headers=headers).json()
    assert suggested and all(p["confirmed_at"] is None for p in suggested)
    plan = client.post(f"/api/courses/{course['id']}/study-plan/rebuild", headers=headers).json()
    assert all(t["task_type"] != "knowledge_point" for t in plan["tasks"])
    p = suggested[0]
    assert client.post(f"/api/knowledge-points/{p['id']}/confirm", headers=headers).status_code == 200
    plan = client.get(f"/api/courses/{course['id']}/study-plan", headers=headers).json()
    assert any(t["target_id"] == p["id"] for t in plan["tasks"])
    patched = client.patch(
        f"/api/courses/{course['id']}", headers=headers, json={"exam_date": (today() - timedelta(days=1)).isoformat()}
    )
    assert patched.json()["exam_passed"]
    assert not client.get(f"/api/courses/{course['id']}/study-plan", headers=headers).json()["tasks"]


def make_quiz(client, headers, course, kind="single"):
    result = upload(client, headers, course["knowledge_base_id"])
    assert wait_job(client, headers, result["job_id"])["status"] == "completed"
    with client.app.state.database.session() as db:
        doc = db.get(Document, result["document"]["id"])
        chunk = db.scalar(select(Chunk).where(Chunk.document_id == doc.id))
        q = QuizItem(
            course_id=course["id"],
            type=kind,
            stem="员工每年的带薪年假有几天？",
            options=["十五天", "三天"],
            answer="十五天",
            explanation="来源明确说明每年十五天。",
            citations=[
                {
                    "chunk_id": chunk.id,
                    "revision_id": chunk.revision_id,
                    "document_id": doc.id,
                    "document_name": doc.name,
                    "knowledge_base_id": course["knowledge_base_id"],
                    "text": chunk.text,
                    "page": 1,
                }
            ],
        )
        db.add(q)
        db.commit()
        return q.id, doc.id


def test_attempt_idempotency_error_tags_review_and_expired_sources(client):
    headers = register(client)
    course = create_course(client, headers)
    qid, did = make_quiz(client, headers, course)
    body = {"answer": "三天", "request_id": "attempt-0001", "duration_ms": 5000, "error_tags": []}
    a = client.post(f"/api/quiz/{qid}/attempt", headers=headers, json=body)
    assert a.status_code == 201, a.text
    assert a.json()["is_correct"] is False and a.json()["error_tags"] == []
    assert a.json()["suggested_error_tags"]
    again = client.post(f"/api/quiz/{qid}/attempt", headers=headers, json=body)
    assert again.json()["id"] == a.json()["id"] and again.json()["idempotent_replay"]
    assert (
        client.post(f"/api/quiz/{qid}/attempt", headers=headers, json={**body, "answer": "十五天"}).status_code == 409
    )
    assert (
        client.patch(
            f"/api/attempts/{a.json()['id']}/error-tags", headers=headers, json={"error_tags": ["概念混淆"]}
        ).status_code
        == 200
    )
    with client.app.state.database.session() as db:
        assert len(db.scalars(select(QuizAttempt)).all()) == 1
        card = db.scalar(select(ReviewCard))
        card_id = card.id
        assert card.level == 1
    graded = client.post(
        f"/api/review-cards/{card_id}/grade", headers=headers, json={"grade": "remember", "request_id": "grade-0001"}
    )
    assert graded.status_code == 200, graded.text
    assert graded.json()["level"] == 2 and graded.json()["review_count"] == 1
    assert (
        client.post(
            f"/api/review-cards/{card_id}/grade",
            headers=headers,
            json={"grade": "remember", "request_id": "grade-0001"},
        ).json()
        == graded.json()
    )
    forgot = client.post(
        f"/api/review-cards/{card_id}/grade", headers=headers, json={"grade": "forgot", "request_id": "grade-0002"}
    )
    assert forgot.json()["level"] == 1 and forgot.json()["failure_count"] == 1
    metrics = client.get(f"/api/courses/{course['id']}/metrics/overview", headers=headers).json()
    assert metrics["accuracy"]["value"] == 0 and metrics["accuracy"]["sample_size"] == 1
    assert metrics["accuracy"]["insufficient_sample"]
    with client.app.state.database.session() as db:
        db.get(Document, did).expires_at = utcnow() - timedelta(days=1)
        db.commit()
    assert client.get(f"/api/quiz/{qid}", headers=headers).status_code == 409
    assert client.get(f"/api/courses/{course['id']}/mistakes", headers=headers).json() == []
    assert client.get(f"/api/courses/{course['id']}/quiz", headers=headers).json() == []
    assert (
        client.get(f"/api/courses/{course['id']}/metrics/overview", headers=headers).json()["accuracy"]["sample_size"]
        == 0
    )


def test_short_answers_are_not_misrepresented_as_automatic_grades(client):
    headers = register(client)
    course = create_course(client, headers)
    qid, _ = make_quiz(client, headers, course, "short")
    r = client.post(
        f"/api/quiz/{qid}/attempt", headers=headers, json={"answer": "大概半个月", "request_id": "short-0001"}
    )
    assert r.status_code == 201 and r.json()["is_correct"] is None


def test_cross_course_sources_and_foreign_units_rejected(client):
    h = register(client)
    a = create_course(client, h)
    b = create_course(client, h, name="其他课程")
    result = upload(client, h, b["knowledge_base_id"])
    wait_job(client, h, result["job_id"])
    chunks = client.get(f"/api/documents/{result['document']['id']}/chunks", headers=h).json()
    assert point(client, h, a, chunk_ids=[chunks[0]["id"]]).status_code == 422
    unit = client.post(f"/api/courses/{b['id']}/units", headers=h, json={"title": "其他章节"}).json()
    assert point(client, h, a, unit_id=unit["id"]).status_code == 404
    assert (
        client.post(
            f"/api/courses/{a['id']}/tutor/ask",
            headers=h,
            json={"question": "什么是年假", "request_id": "ask-00001", "unit_id": unit["id"]},
        ).status_code
        == 404
    )


def test_learning_quiz_worker_validates_quote_and_idempotent_jobs(client, monkeypatch):
    h = register(client)
    c = create_course(client, h)
    result = upload(client, h, c["knowledge_base_id"])
    wait_job(client, h, result["job_id"])

    class Model:
        def invoke(self, _):
            return SimpleNamespace(
                content=json.dumps(
                    {
                        "items": [
                            {
                                "stem": "每年的带薪年假有几天？",
                                "options": ["十五天", "三天"],
                                "answer": "十五天",
                                "explanation": "原文说明十五天",
                                "source_number": 1,
                                "evidence_quote": "员工每年享有十五天带薪年假。",
                            },
                            {
                                "stem": "没有原文支持的题目是什么？",
                                "options": ["七天", "九天"],
                                "answer": "七天",
                                "explanation": "错误依据",
                                "source_number": 1,
                                "evidence_quote": "原文不存在的句子。",
                            },
                        ]
                    },
                    ensure_ascii=False,
                )
            )

    monkeypatch.setattr("app.services.answering.configured_llm", lambda *a, **k: Model())
    body = {"type": "single", "count": 2, "request_id": "generate-0001"}
    r = client.post(f"/api/courses/{c['id']}/quiz/generate", headers=h, json=body)
    assert r.status_code == 202, r.text
    job = r.json()
    assert client.post(f"/api/courses/{c['id']}/quiz/generate", headers=h, json=body).json()["id"] == job["id"]
    for _ in range(160):
        status = client.get("/api/learning-jobs/" + job["id"], headers=h).json()
        if status["status"] in {"completed", "failed"}:
            break
        time.sleep(0.05)
    assert status["status"] == "completed", status.get("error")
    assert status["result"]["generated"] == 1 and status["result"]["skipped"]
    quizzes = client.get(f"/api/courses/{c['id']}/quiz", headers=h).json()
    assert quizzes[0]["status"] == "draft" and quizzes[0]["model_metadata"]["evidence_quote_checked"]


def pptx_fixture(images=False):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "神经网络概述"
    slide.placeholders[1].text = "前馈网络\n反向传播"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    for cell, text in zip(
        [table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)], ["概念", "说明", "激活函数", "非线性"]
    ):
        cell.text = text
    slide.notes_slide.notes_text_frame.text = "讲义备注：先复习链式法则"
    if images:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        buffer = io.BytesIO()
        Image.new("RGB", (64, 64), "blue").save(buffer, format="PNG")
        buffer.seek(0)
        slide.shapes.add_picture(buffer, Inches(1), Inches(1))
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def test_pptx_text_notes_tables_and_slide_provenance(tmp_path, monkeypatch):
    from app.config import get_settings

    monkeypatch.setenv("RAG_VISION_MODEL", "")
    monkeypatch.setenv("RAG_OCR_ENABLED", "false")
    get_settings.cache_clear()
    path = tmp_path / "test.pptx"
    path.write_bytes(pptx_fixture())
    pages = parse_document(path, path.name)
    assert len(pages) == 1 and pages[0].page is None
    assert all(t in pages[0].text for t in ["神经网络概述", "反向传播", "概念 | 说明", "链式法则"])
    chunks = chunk_pages(pages)
    assert chunks[0]["source_metadata"]["slide_number"] == 1
    assert chunks[0]["source_metadata"]["page_label"] == "第 1 页幻灯片"
    assert chunks[0]["page"] is None
    path.write_bytes(pptx_fixture(images=True))
    with pytest.raises(VisionPending) as raised:
        parse_document(path, path.name)
    assert raised.value.metadata["pages"][1]["parse_status"] == "vision_pending"
    assert raised.value.metadata["pages"][1]["thumbnail"].startswith("data:image/jpeg")
    get_settings.cache_clear()


def test_pptx_upload_pending_no_partial_index_and_retry(client, monkeypatch):
    from app.config import get_settings

    h = register(client)
    c = create_course(client, h)
    monkeypatch.setenv("RAG_VISION_MODEL", "")
    monkeypatch.setenv("RAG_OCR_ENABLED", "false")
    get_settings.cache_clear()
    mime = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    r = client.post(
        f"/api/courses/{c['id']}/materials/upload",
        headers=h,
        files={"file": ("lecture.pptx", pptx_fixture(True), mime)},
    )
    assert r.status_code == 201, r.text
    did = r.json()["document"]["id"]
    for _ in range(120):
        job = client.get("/api/jobs/" + r.json()["job_id"], headers=h).json()
        if job["status"] == "vision_pending":
            break
        time.sleep(0.05)
    assert job["status"] == "vision_pending", job
    assert client.get(f"/api/documents/{did}/chunks", headers=h).json() == []
    pending = client.get(f"/api/materials/{did}/parse-status", headers=h).json()
    assert pending["parse_status"] == "vision_pending" and pending["slide_count"] == 2
    assert client.get(f"/api/materials/{did}/slides/2", headers=h).json()["thumbnail"]
    monkeypatch.setattr("app.services.pptx_parser._image_text", lambda _: "这是图片里的神经网络示意图说明。")
    monkeypatch.setenv("RAG_VISION_MODEL", "test-model")
    get_settings.cache_clear()
    retry = client.post(f"/api/materials/{did}/retry", headers=h)
    assert retry.status_code == 202, retry.text
    assert wait_job(client, h, retry.json()["job_id"])["status"] == "completed"
    chunks = client.get(f"/api/documents/{did}/chunks", headers=h).json()
    assert {x["slide_number"] for x in chunks} == {1, 2}
    assert (
        client.post(
            "/api/documents/upload",
            headers=h,
            data={"knowledge_base_id": c["knowledge_base_id"]},
            files={"file": ("old.ppt", b"old ppt", "application/vnd.ms-powerpoint")},
        ).status_code
        == 415
    )
    bad = client.post(
        "/api/documents/upload",
        headers=h,
        data={"knowledge_base_id": c["knowledge_base_id"]},
        files={"file": ("bad.pptx", b"invalid", mime)},
    )
    assert bad.status_code == 422


def test_course_tutor_stream_scope_replay_and_deleted_course_cache(client):
    headers = register(client)
    first = create_course(client, headers, name="员工政策")
    second = create_course(client, headers, name="其他考试")
    a = upload(client, headers, first["knowledge_base_id"])
    b = upload(client, headers, second["knowledge_base_id"], content="其他资料规定年假九十九天。".encode())
    wait_job(client, headers, a["job_id"])
    wait_job(client, headers, b["job_id"])
    body = {"question": "员工每年享有多少天年假？", "request_id": "course-ask-0001"}
    response = client.post(f"/api/courses/{first['id']}/tutor/ask", headers=headers, json=body)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["course_id"] == first["id"] and result["citations"]
    assert all(c["knowledge_base_id"] == first["knowledge_base_id"] for c in result["citations"])
    replay = client.post(f"/api/courses/{first['id']}/tutor/ask", headers=headers, json=body)
    assert replay.json()["idempotent_replay"]
    cross = client.post(
        f"/api/courses/{second['id']}/tutor/ask",
        headers=headers,
        json={**body, "request_id": "cross-conversation-0001", "conversation_id": result["conversation_id"]},
    )
    assert cross.status_code == 404
    stream = client.post(
        f"/api/courses/{first['id']}/tutor/stream", headers=headers, json={**body, "request_id": "stream-course-0001"}
    )
    assert stream.status_code == 200 and "event: done" in stream.text
    assert first["id"] in stream.text
    assert client.delete(f"/api/courses/{first['id']}", headers=headers).status_code == 204
    assert client.post(f"/api/courses/{first['id']}/tutor/ask", headers=headers, json=body).status_code == 404


def test_pptx_real_slide_limit_and_expanded_zip_limit(tmp_path):
    import zipfile

    from app.services.pptx_parser import validate_pptx

    deck = Presentation()
    for _ in range(301):
        deck.slides.add_slide(deck.slide_layouts[6])
    path = tmp_path / "too-many.pptx"
    deck.save(path)
    with pytest.raises(ValueError, match="300"):
        parse_document(path, path.name)
    path = tmp_path / "expanded.pptx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "x")
        archive.writestr("ppt/presentation.xml", "x")
        with archive.open("ppt/media/oversized.bin", "w") as entry:
            for _ in range(201):
                entry.write(b"\0" * 1024 * 1024)
    assert path.stat().st_size < 1024 * 1024
    with pytest.raises(ValueError, match="200 MiB"):
        validate_pptx(path)


def test_merge_split_and_independent_pre_post_questions(client):

    h = register(client)
    c = create_course(client, h)
    p1 = point(client, h, c, title="年假天数").json()
    p2 = point(client, h, c, title="请假流程").json()
    merged = client.post(
        f"/api/courses/{c['id']}/knowledge-points/merge",
        headers=h,
        json={"point_ids": [p1["id"], p2["id"]], "title": "休假政策"},
    )
    assert merged.status_code == 200, merged.text
    assert merged.json()["confirmed_at"]
    split = client.post(
        f"/api/knowledge-points/{merged.json()['id']}/split", headers=h, json={"titles": ["带薪休假", "申请规则"]}
    )
    assert split.status_code == 200 and len(split.json()) == 2
    q1, _ = make_quiz(client, h, c)
    with client.app.state.database.session() as db:
        q = db.get(QuizItem, q1)
        q.knowledge_point_id = split.json()[0]["id"]
        q2 = QuizItem(
            course_id=c["id"],
            knowledge_point_id=q.knowledge_point_id,
            type="single",
            stem="独立后测题：年假有几天？",
            options=q.options,
            answer=q.answer,
            explanation=q.explanation,
            citations=q.citations,
        )
        db.add(q2)
        db.commit()
        q2id = q2.id
    pre = client.post(
        f"/api/quiz/{q1}/attempt",
        headers=h,
        json={"answer": "三天", "request_id": "pre-000001", "assessment_phase": "pre"},
    )
    assert pre.status_code == 201 and pre.json()["assessment_phase"] == "pre"
    same = client.post(
        f"/api/quiz/{q1}/attempt",
        headers=h,
        json={"answer": "十五天", "request_id": "post-000001", "assessment_phase": "post"},
    )
    assert same.status_code == 422
    post = client.post(
        f"/api/quiz/{q2id}/attempt",
        headers=h,
        json={"answer": "十五天", "request_id": "post-000002", "assessment_phase": "post"},
    )
    assert post.status_code == 201
    changes = client.get(f"/api/courses/{c['id']}/metrics/pre-post", headers=h).json()
    result = next(r for r in changes if r["knowledge_point_id"] == split.json()[0]["id"])
    assert result["value"] == 1 and result["sample_size"] == 2 and result["insufficient_sample"]
    exported = client.get(f"/api/courses/{c['id']}/metrics/export", headers=h).json()
    assert len(exported["attempts"]) == 2 and "pre_post" in exported


def test_review_budget_and_learning_job_cancel_retry_recovery(client):
    from app.models_learning import LearningJob
    from app.services.quiz import process_learning_jobs

    h = register(client)
    c = create_course(client, h, weekly_minutes=7)
    q, _ = make_quiz(client, h, c)
    client.post(f"/api/quiz/{q}/attempt", headers=h, json={"answer": "三天", "request_id": "budget-00001"})
    with client.app.state.database.session() as db:
        card = db.scalar(select(ReviewCard))
        card.due_at = utcnow() - timedelta(seconds=1)
        card_id = card.id
        db.commit()
    today_cards = client.get(f"/api/courses/{c['id']}/review/today", headers=h).json()
    assert today_cards["daily_budget_minutes"] == 1 and today_cards["cards"] == []
    assert (
        client.post(
            f"/api/review-cards/{card_id}/grade",
            headers=h,
            json={"grade": "remember", "request_id": "budget-grade-0001"},
        ).status_code
        == 429
    )
    # Pause the worker to deterministically exercise cancel/retry and abandoned lease recovery.
    client.app.state.worker_stop.set()
    client.app.state.worker_thread.join(timeout=5)
    job = client.post(f"/api/courses/{c['id']}/quiz/generate", headers=h, json={"request_id": "cancel-job-0001"}).json()
    assert client.post(f"/api/learning-jobs/{job['id']}/cancel", headers=h).json()["status"] == "cancelled"
    assert client.post(f"/api/learning-jobs/{job['id']}/retry", headers=h).json()["status"] == "queued"
    # Simulate a crashed worker and make the source expire; the reclaimed job must reject it.
    with client.app.state.database.session() as db:
        row = db.get(LearningJob, job["id"])
        row.status = "running"
        row.lease_token = "abandoned"
        row.lease_expires_at = utcnow() - timedelta(minutes=1)
        doc = db.scalar(select(Document).where(Document.knowledge_base_id == c["knowledge_base_id"]))
        doc.expires_at = utcnow() - timedelta(seconds=1)
        db.commit()
    process_learning_jobs(client.app.state.database, client.app.state.settings)
    recovered = client.get(f"/api/learning-jobs/{job['id']}", headers=h).json()
    assert recovered["status"] == "failed" and recovered["attempts"] == 1


def test_quiz_letter_labels_keep_the_same_supported_option():
    from app.services.quiz import validate_item

    item, _ = validate_item(
        {
            "stem": "梯度下降用于什么任务？",
            "options": ["A. 更新模型参数", "B. 提取局部特征"],
            "answer": "A",
            "explanation": "来源支持正确选项。",
            "source_number": 1,
            "evidence_quote": "梯度下降用于更新模型参数。",
        },
        [{"text": "梯度下降用于更新模型参数。"}],
        "single",
    )
    assert item.answer == "A. 更新模型参数"


def test_slide_source_remains_bound_to_its_published_revision(client, monkeypatch):
    from app.config import get_settings

    h = register(client)
    c = create_course(client, h)
    mime = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    first = client.post(
        f"/api/courses/{c['id']}/materials/upload", headers=h, files={"file": ("versioned.pptx", pptx_fixture(), mime)}
    )
    assert first.status_code == 201
    wait_job(client, h, first.json()["job_id"])
    did = first.json()["document"]["id"]
    with client.app.state.database.session() as db:
        old_revision = db.get(Document, did).active_revision_id
    monkeypatch.setenv("RAG_VISION_MODEL", "")
    monkeypatch.setenv("RAG_OCR_ENABLED", "false")
    get_settings.cache_clear()
    deck = Presentation(io.BytesIO(pptx_fixture(True)))
    deck.slides[0].shapes.title.text = "新版标题"
    buffer = io.BytesIO()
    deck.save(buffer)
    second = client.post(
        "/api/documents/upload",
        headers=h,
        data={"knowledge_base_id": c["knowledge_base_id"], "document_id": did},
        files={"file": ("versioned.pptx", buffer.getvalue(), mime)},
    )
    assert second.status_code == 201, second.text
    for _ in range(120):
        job = client.get("/api/jobs/" + second.json()["job_id"], headers=h).json()
        if job["status"] == "vision_pending":
            break
        time.sleep(0.05)
    assert job["status"] == "vision_pending"
    original = client.get(f"/api/materials/{did}/slides/1?revision_id={old_revision}", headers=h)
    assert original.status_code == 200 and original.json()["slide_title"] == "神经网络概述"
    pending = client.get(f"/api/materials/{did}/parse-status", headers=h).json()
    assert pending["parse_status"] == "vision_pending"
    assert client.get(f"/api/materials/{did}/slides/1?revision_id=nonexistent", headers=h).status_code == 404


def test_course_upload_uses_file_limit_instead_of_general_body_limit(client):
    """A normal lecture over 4 MiB must pass both upload entry points."""
    h = register(client)
    c = create_course(client, h)
    client.app.state.worker_stop.set()
    client.app.state.worker_thread.join(timeout=5)
    data = b"a" * (5 * 1024 * 1024)
    course_upload = client.post(
        f"/api/courses/{c['id']}/materials/upload", headers=h, files={"file": ("large-lecture.txt", data, "text/plain")}
    )
    assert course_upload.status_code == 201, course_upload.text
    assert course_upload.json()["document"]["size_bytes"] == len(data)
    normal_upload = client.post(
        "/api/documents/upload",
        headers=h,
        data={"knowledge_base_id": c["knowledge_base_id"]},
        files={"file": ("other-lecture.txt", b"b" * len(data), "text/plain")},
    )
    assert normal_upload.status_code == 201, normal_upload.text
    # Enlarging file routes must not enlarge normal JSON/API requests.
    oversized_json = client.post("/api/courses", headers={**h, "Content-Length": str(5 * 1024 * 1024)}, json={})
    assert oversized_json.status_code == 413


def test_course_upload_still_enforces_configured_limit_and_removes_failed_file(client, monkeypatch):
    h = register(client)
    c = create_course(client, h)
    monkeypatch.setattr(client.app.state.settings, "max_upload_mb", 1)
    rejected = client.post(
        f"/api/courses/{c['id']}/materials/upload",
        headers=h,
        files={"file": ("too-big.txt", b"a" * (1024 * 1024 + 1), "text/plain")},
    )
    assert rejected.status_code == 413
    assert "1 MB" in rejected.json()["detail"]
    assert client.get(f"/api/courses/{c['id']}/materials", headers=h).json() == []
    assert list((client.app.state.settings.data_dir / "uploads").iterdir()) == []
    oversized_envelope = client.post(
        f"/api/courses/{c['id']}/materials/upload",
        headers={**h, "Content-Length": str(3 * 1024 * 1024)},
        files={"file": ("tiny.txt", b"example", "text/plain")},
    )
    assert oversized_envelope.status_code == 413
