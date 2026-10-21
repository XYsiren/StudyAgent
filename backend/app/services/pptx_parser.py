"""Bounded OOXML parsing: slide identity never reuses PDF page numbers."""

import base64
import io
import zipfile

from app.config import get_settings
from app.services.parsing import ParsedPage, _image_text

MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


class VisionPending(ValueError):
    def __init__(self, metadata):
        super().__init__("部分幻灯片需要视觉解析，请配置视觉模型后重试；未发布部分索引")
        self.metadata = metadata


def validate_pptx(path):
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > 20000 or sum(i.file_size for i in infos) > 200 * 1024 * 1024:
                raise ValueError("PPTX 解压大小超过 200 MiB 或对象数量超限")
            names = {i.filename for i in infos}
            if "[Content_Types].xml" not in names or "ppt/presentation.xml" not in names:
                raise ValueError("文件不是有效的 PPTX 演示文稿")
            if any("vbaProject" in name for name in names):
                raise ValueError("不接受包含宏的演示文稿，请另存为无宏 PPTX")
    except zipfile.BadZipFile as exc:
        raise ValueError("PPTX 文件已损坏，请重新导出为 .pptx") from exc


def parse_pptx(path, progress=None):
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    validate_pptx(path)
    try:
        presentation = Presentation(path)
    except Exception as exc:
        raise ValueError("PPTX 无法解析，请重新导出为 .pptx") from exc
    total = len(presentation.slides)
    if total > 300:
        raise ValueError("PPTX 超过 300 页幻灯片，请拆分后上传")
    pages, pending = [], False
    settings = get_settings()
    for number, slide in enumerate(presentation.slides, 1):
        if progress:
            progress(number - 1, total, "提取幻灯片")
        lines, images, unparsed = [], [], 0
        flags = set()
        has_readable_text = False
        title = slide.shapes.title.text if slide.shapes.title else ""

        def walk(shapes):
            nonlocal unparsed, has_readable_text
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(shape.shapes)
                    continue
                alt = shape._element.xpath(".//p:cNvPr")
                for node in alt:
                    description = node.get("descr") or node.get("title")
                    if description:
                        lines.append("[替代文本] " + description)
                if shape.has_text_frame:
                    has_readable_text = has_readable_text or bool(shape.text.strip())
                    lines.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                    flags.add("text")
                elif shape.has_table:
                    has_readable_text = True
                    for row in shape.table.rows:
                        lines.append(" | ".join(cell.text for cell in row.cells))
                    flags.add("table")
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    images.append(shape.image.blob)
                    flags.add("image")
                elif shape.has_chart:
                    flags.add("image")
                    unparsed += 1
                    chart = shape.chart
                    if chart.has_title:
                        lines.append("[图表] " + chart.chart_title.text_frame.text)
                    for series in chart.series:
                        lines.append(f"{series.name}: " + ", ".join(str(v) for v in series.values))
                else:
                    unparsed += 1

        walk(slide.shapes)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes and notes.text.strip():
                lines.append("[备注]\n" + notes.text)
        state, thumbnail = "ready", None
        if images and not has_readable_text:
            if settings.ocr_enabled or settings.vision_model:
                if progress:
                    progress(number - 1, total, "OCR / 视觉识别")
                for blob in images:
                    lines.append(_image_text(blob))
            else:
                pending = True
                state = "vision_pending"
                unparsed += len(images)
                with Image.open(io.BytesIO(images[0])) as image:
                    image.thumbnail((360, 240))
                    buffer = io.BytesIO()
                    image.convert("RGB").save(buffer, format="JPEG", quality=65)
                    thumbnail = "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode()
        elif images:
            # Textual slides remain searchable; image omissions stay visible in quality metadata.
            unparsed += len(images)
        metadata = {
            "source_type": "pptx",
            "slide_number": number,
            "page_label": f"第 {number} 页幻灯片",
            "slide_title": title,
            "page_type": next(iter(flags)) if len(flags) == 1 else "mixed",
            "unparsed_objects": unparsed,
            "parse_status": state,
        }
        if thumbnail:
            metadata["thumbnail"] = thumbnail
        pages.append(ParsedPage("\n".join(lines).strip(), None, metadata))
        if progress:
            progress(number, total, "幻灯片提取完成")
    if pending:
        raise VisionPending(
            {
                "file_type": "pptx",
                "slide_count": total,
                "pages": [{"text": p.text[:12000], **p.metadata} for p in pages],
                "unparsed_objects": sum(p.metadata["unparsed_objects"] for p in pages),
            }
        )
    return pages
